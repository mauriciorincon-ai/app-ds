"""S18 (ADR 025): el kit de prueba en Excel.

Convierte los CSV de ejemplo del kit a `.xlsx` con openpyxl (un escritor de terceros: la prueba de que el lector de
Probeta abre lo que otro programa escribe, no solo lo que escribe Probeta) y arma los libros de borde de
`tests/fixtures/xlsx/`. No corre en la CI ni en `pnpm test`: los archivos que genera están en el repo y un test los
compara con sus CSV.

    python3 -m venv .venv && .venv/bin/pip install openpyxl==3.1.5
    .venv/bin/python scripts/make-kit-xlsx.py

Los números van como números, las fechas ISO como fechas reales (formato yyyy-mm-dd) y las celdas vacías como celdas
vacías: lo que haría Excel al abrir el CSV y guardarlo.
"""

import csv
import datetime as dt
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

RAIZ = Path(__file__).resolve().parent.parent
KIT = RAIZ / "docs" / "kit-de-prueba"
SALIDA = KIT / "xlsx"
BORDES = RAIZ / "tests" / "fixtures" / "xlsx"

# Los ejemplos que la pantalla de inicio ofrece (public/datasets) y el kit con fuga plantada.
EJEMPLOS = [
    "clientes-sucio",
    "consumo-energia",
    "credito-fuga-plantada",
    "marketing-campania",
    "planes-fuga-plantada",
    "planes-suscripcion",
    "precio-fuga-plantada",
    "rotacion-empleados",
    "segmentos-clientes",
    "sin-grupos",
    "clientes-con-datos-personales",
]

FECHA_FIJA = dt.datetime(2026, 10, 25, 12, 0, 0)
ENTERO = re.compile(r"^-?(0|[1-9]\d*)$")
DECIMAL = re.compile(r"^-?(0|[1-9]\d*)\.\d+$")
FECHA = re.compile(r"^(\d{4})-(\d{2})-(\d{2})$")


def celda(texto):
    """El valor que Excel habría puesto en la celda al abrir el CSV."""
    if texto == "":
        return None
    if ENTERO.match(texto) and len(texto) <= 15:
        return int(texto)
    if DECIMAL.match(texto):
        return float(texto)
    f = FECHA.match(texto)
    if f:
        try:
            return dt.date(int(f[1]), int(f[2]), int(f[3]))
        except ValueError:
            return texto
    return texto


def libro():
    wb = Workbook()
    wb.properties.created = FECHA_FIJA
    wb.properties.modified = FECHA_FIJA
    wb.properties.creator = "Probeta DS · kit de prueba"
    return wb


def del_csv(nombre):
    with open(KIT / f"{nombre}.csv", encoding="utf-8-sig", newline="") as f:
        filas = list(csv.reader(f))
    wb = libro()
    ws = wb.active
    ws.title = "Datos"
    for i, fila in enumerate(filas):
        for j, texto in enumerate(fila):
            valor = celda(texto) if i > 0 else texto
            if valor is None:
                continue
            c = ws.cell(row=i + 1, column=j + 1, value=valor)
            if isinstance(valor, dt.date):
                c.number_format = "yyyy-mm-dd"
    wb.save(SALIDA / f"{nombre}.xlsx")


def bordes():
    BORDES.mkdir(parents=True, exist_ok=True)

    # 1) tipos mezclados, hoja oculta primero y una segunda hoja que se ignora
    wb = libro()
    oculta = wb.active
    oculta.title = "Oculta"
    oculta.sheet_state = "hidden"
    oculta.append(["no", "debe", "leerse"])
    ws = wb.create_sheet("Datos")
    ws.append(["nombre", "monto", "activo", "alta", "visita", "hora", "nota"])
    ws.append(["Ana López", 1200.5, True, dt.date(2026, 1, 15), dt.datetime(2026, 2, 3, 14, 30, 0), dt.time(8, 5, 0), "a & b <c>"])
    ws.append(["Luis", 0.1, False, dt.date(2026, 12, 31), dt.datetime(2026, 2, 3, 0, 0, 0), None, "línea 1\nlínea 2"])
    ws.append([None, 3, None, None, None, None, None])  # fila con un solo valor
    ws.append([])  # fila vacía
    ws.append(["Zoë", 1e-7, True, dt.date(1900, 3, 1), None, None, "ñandú"])
    ws["H1"] = "dispersa"
    ws["H5"] = 7  # columna que solo tiene un valor, lejos de las demás
    for c in ("D2", "D3", "D5"):
        ws[c].number_format = "yyyy-mm-dd"
    ws["E2"].number_format = "yyyy-mm-dd hh:mm:ss"
    ws["E3"].number_format = "yyyy-mm-dd hh:mm:ss"
    ws["F2"].number_format = "hh:mm:ss"
    otra = wb.create_sheet("Otra")
    otra.append(["x"])
    wb.save(BORDES / "bordes.xlsx")

    # 2) sistema de fechas de 1904 (Excel para Mac antiguo)
    wb = libro()
    wb.epoch = CALENDAR_MAC_1904
    ws = wb.active
    ws.title = "Fechas"
    ws.append(["dia"])
    ws.append([dt.date(2026, 1, 15)])
    ws.append([dt.date(2030, 7, 4)])
    ws["A2"].number_format = "yyyy-mm-dd"
    ws["A3"].number_format = "yyyy-mm-dd"
    wb.save(BORDES / "fechas-1904.xlsx")


def malos():
    """Los tres libros que Probeta tiene que rechazar diciendo por qué (los mismos bytes del e2e `excel-y-limpieza`).

    No son libros de verdad: llevan solo la firma que el lector reconoce, y ninguna celda de nadie.
    """
    ole = bytes([0xD0, 0xCF, 0x11, 0xE0, 0xA1, 0xB1, 0x1A, 0xE1])

    def utf16(texto):
        return b"".join(bytes([ord(c), 0]) for c in texto)

    (SALIDA / "libro-con-clave.xlsx").write_bytes(ole + bytes(200) + utf16("EncryptedPackage"))
    (SALIDA / "libro-antiguo.xls").write_bytes(ole + utf16("Workbook") + bytes(200))
    (SALIDA / "libro-danado.xlsx").write_bytes(
        ("no es un libro de verdad, solo texto suelto" * 3).encode("utf-8")
    )


def main():
    """Sin argumentos rehace todo; con nombres (sin extensión) rehace solo esos ejemplos."""
    SALIDA.mkdir(parents=True, exist_ok=True)
    nombres = sys.argv[1:] or EJEMPLOS
    for nombre in nombres:
        if (KIT / f"{nombre}.csv").exists():
            del_csv(nombre)
    if not sys.argv[1:]:
        bordes()
        malos()
    print("kit en", SALIDA, "· bordes en", BORDES)


if __name__ == "__main__":
    main()
